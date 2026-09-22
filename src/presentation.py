"""
presentation.py
===============
Build a reader-friendly Excel workbook from the consolidated long-format data.
For PRESENTATION / interpretation, not data management.

LAYOUT (one worksheet per scale)
--------------------------------
Publication by publication, stacked down the sheet. For EACH publication:

    <publication name>
    sample_type   healthy_controls              | patients
    sample_size subsample EA EN ... total score | sample_size subsample EA ...
    <one row per healthy_controls subsample>    | <one row per patients subsample>
    total   <weighted over that pub's subs>     | total  <weighted ...>

After all publications, a GRAND TOTAL section aggregates over every publication.
healthy_controls sits on the LEFT, patients on the RIGHT, separated by a blank
column. If a publication has no data for one sample_type, that side stays blank.

WEIGHTING ("Option B")
----------------------
Every value is a sample-size-weighted mean of the underlying non-redundant
means, computed with the project's aggregate.aggregate():
  - subsample cell : that subsample's weighted mean
  - publication "total" : weighted over that publication's subsamples
  - "grand total" : weighted over all publications' subsamples
Redundant rows (redundant_aggregate == True) are excluded by aggregate();
only data_type == "mean" is shown.

WHY EXPLICIT CELL COORDINATES
-----------------------------
The two side-by-side blocks can have different numbers of rows, and either can
be empty. Writing whole dataframes with to_excel() shifts columns in those
cases, which would misalign values. So instead we compute each block's values
as a grid and write every cell to an explicit (row, column) address with
openpyxl. A value can then never land in a column it does not belong to, and an
empty block simply has no cells written.

COLUMN ORDER
------------
`subscale_orders` fixes column order per scale (e.g. keep EA next to EN).
Listed subscales first, others appended alphabetically, "total score" last.

SAMPLE SIZE COLUMNS: single vs. per-measure
--------------------------------------------
Normally one subsample answers every subscale, so one shared "sample_size"
column per row is unambiguous. Some publications (see the subsample-naming
addendum) report a DIFFERENT n per subscale within the same block -- e.g.
PSYRATS's AHS and DS scores come from different, overlapping counts of
participants. A single "sample_size" number on a "total" row that reports
several subscales at once would then be meaningless (it would blend two
unrelated counts).

Which publications have this problem is a JUDGEMENT CALL, not something
this module infers from the numbers: a subscale's summed n can differ from
another subscale's for a completely innocent reason too (e.g. a subsample
simply wasn't reported for that subscale -- CTQ_SF has plenty of these), so
comparing sums alone gives false positives. The caller passes in
`ambiguous_publications`, a plain list/set of publication names known (from
your own eyeballing at extraction time) to have this overlap problem.

This is decided ONCE PER SHEET (i.e. per scale, since column headers must be
identical down the whole sheet): if none of a scale's publications are in
`ambiguous_publications`, the sheet keeps the classic single "sample_size"
column. If any are, the WHOLE sheet switches to one "n_<measure>" column
immediately to the left of each measure column instead (subsample, n_AHS,
AHS, n_DS, DS, n_total, total score, ...) -- every publication on that sheet
then uses this wider layout, even ones whose own n is uniform, so the grid
stays one consistent shape. See `_needs_expanded_columns`.

MAIN FUNCTION
-------------
  write_presentation_excel(df, output_path, subscale_orders=None,
                            ambiguous_publications=None)
"""

import pandas as pd
from openpyxl import Workbook

import aggregate  # the project's existing weighting module

TOTAL_COLUMN_LABEL = "total score"
# Fixed left-to-right order of the two blocks.
BLOCK_ORDER = ["healthy_controls", "patients"]


# ---------------------------------------------------------------------------
# Decide the measure columns (subscales + total score) for one scale.
# ---------------------------------------------------------------------------
def _measure_columns(scale_df, subscale_order=None):
    """Return the ordered column names for one scale: subscales then total."""
    present = sorted(v for v in scale_df["subscale"].dropna().unique()
                     if v != "none")
    if subscale_order:
        ordered = [s for s in subscale_order if s in present]
        leftovers = sorted(s for s in present if s not in subscale_order)
        subscales = ordered + leftovers
    else:
        subscales = present
    return subscales + [TOTAL_COLUMN_LABEL]


# ---------------------------------------------------------------------------
# Look up one weighted value from aggregated results.
# ---------------------------------------------------------------------------
def _lookup(agg_df, **conditions):
    """Return the single matching weighted_value, or None if not exactly one."""
    mask = pd.Series(True, index=agg_df.index)
    for col, val in conditions.items():
        mask &= agg_df[col] == val
    hits = agg_df[mask]
    if len(hits) == 1:
        return hits["weighted_value"].iloc[0]
    return None

# ---------------------------------------------------------------------------
# The measure values for one row (one subsample, or a total), as a list
# aligned to `measures`.
# ---------------------------------------------------------------------------

def _measure_values(agg_df, measures, base_conditions):
    """Return a list of values, one per measure column, in order.

    A subscale column matches rows that are BOTH that subscale AND
    record_type == "subscale" — so item-level rows for the same subscale name
    are never picked up. The "total score" column matches record_type == "total".
    """
    values = []
    for m in measures:
        if m == TOTAL_COLUMN_LABEL:
            values.append(_lookup(agg_df, **base_conditions, record_type="total"))
        else:
            values.append(_lookup(agg_df, **base_conditions, subscale=m,
                                  record_type="subscale"))
    return values

# ---------------------------------------------------------------------------
# Sample size for one subsample (single-column / non-expanded layout only).
# ---------------------------------------------------------------------------
def _subsample_size(scale_df, publication, sample_type, subsample):
    """Return the sample_size for one subsample (first numeric n found).

    Only valid when the sheet is NOT using expanded per-measure columns, i.e.
    this subsample's n is the same regardless of which measure it's reporting.
    """
    rows = scale_df[
        (scale_df["publication"] == publication)
        & (scale_df["sample_type"] == sample_type)
        & (scale_df["subsample"] == subsample)
    ]
    sizes = pd.to_numeric(rows["sample_size"], errors="coerce").dropna()
    return int(sizes.iloc[0]) if len(sizes) else None


# ---------------------------------------------------------------------------
# Sample size for one (subsample, measure) cell -- expanded layout.
# ---------------------------------------------------------------------------
def _measure_sample_size(scale_df, publication, sample_type, subsample, measure):
    """Return the n behind ONE subsample's ONE measure cell, or None.

    Mirrors `_measure_values`'s own filtering: the "total score" measure
    looks at record_type=="total"; any other measure looks at that exact
    subscale with record_type=="subscale". This is what lets a single
    subsample row show a different, correct n under each measure column.
    """
    rows = scale_df[
        (scale_df["publication"] == publication)
        & (scale_df["sample_type"] == sample_type)
        & (scale_df["subsample"] == subsample)
    ]
    if measure == TOTAL_COLUMN_LABEL:
        rows = rows[rows["record_type"] == "total"]
    else:
        rows = rows[(rows["subscale"] == measure) & (rows["record_type"] == "subscale")]
    sizes = pd.to_numeric(rows["sample_size"], errors="coerce").dropna()
    return int(sizes.iloc[0]) if len(sizes) else None


# ---------------------------------------------------------------------------
# Summed sample size for a block: each subsample counted ONCE (single column).
# ---------------------------------------------------------------------------
def _summed_sample_size(scale_df, sample_type, publication=None):
    """Return the total n for a block: sum of the distinct subsamples' n.

    A subsample's n is repeated across its mean/sd/subscale rows, so we first
    reduce to one row per subsample (drop_duplicates), then sum. This counts
    each subsample exactly once. If `publication` is None, the sum is over all
    publications (used for the grand total).

    Only non-redundant `mean` rows are considered, matching what is displayed.
    Only valid for a block whose n does NOT vary by subscale/measure -- see
    `_summed_sample_size_per_measure` and `_needs_expanded_columns` for the
    case where it does.
    """
    rows = scale_df[scale_df["sample_type"] == sample_type]
    rows = rows[rows["data_type"] == "mean"]
    if "redundant_aggregate" in rows.columns:
        rows = rows[rows["redundant_aggregate"] != True]  # noqa: E712
    if publication is not None:
        rows = rows[rows["publication"] == publication]

    # one row per subsample (within publication if given), then sum their n
    keys = ["publication", "subsample"]
    distinct = rows.drop_duplicates(keys)
    sizes = pd.to_numeric(distinct["sample_size"], errors="coerce").dropna()
    return int(sizes.sum()) if len(sizes) else None


# ---------------------------------------------------------------------------
# Summed sample size PER MEASURE for a block (expanded layout).
# ---------------------------------------------------------------------------
def _summed_sample_size_per_measure(scale_df, sample_type, measures, publication=None):
    """Return {measure: summed_n} for a block, each subsample counted ONCE
    WITHIN that measure.

    Same idea as `_summed_sample_size`, but grouped by measure (subscale, or
    "total score" via record_type=="total") instead of blending every
    subscale's subsamples together. This is what a "total" row needs when a
    publication's n genuinely differs by subscale (e.g. PSYRATS AHS vs DS) --
    summing across measures would double-count participants who appear under
    more than one subscale.
    """
    rows = scale_df[scale_df["sample_type"] == sample_type]
    rows = rows[rows["data_type"] == "mean"]
    if "redundant_aggregate" in rows.columns:
        rows = rows[rows["redundant_aggregate"] != True]  # noqa: E712
    if publication is not None:
        rows = rows[rows["publication"] == publication]

    result = {}
    for m in measures:
        if m == TOTAL_COLUMN_LABEL:
            m_rows = rows[rows["record_type"] == "total"]
        else:
            m_rows = rows[(rows["subscale"] == m) & (rows["record_type"] == "subscale")]
        distinct = m_rows.drop_duplicates(["publication", "subsample"])
        sizes = pd.to_numeric(distinct["sample_size"], errors="coerce").dropna()
        if len(sizes):
            result[m] = int(sizes.sum())
    return result


# ---------------------------------------------------------------------------
# Decide, once per sheet, whether the sample size needs its own column
# PER MEASURE (rather than one shared "sample_size" column).
# ---------------------------------------------------------------------------
def _needs_expanded_columns(scale_df, ambiguous_publications):
    """True if any publication on this sheet is in `ambiguous_publications`.

    Deliberately NOT inferred from the data (e.g. by comparing summed n's
    across subscales): a subscale's summed n can legitimately differ from
    another's for an innocent reason too -- a subsample simply wasn't
    reported for that subscale (common in CTQ_SF/DERS) -- so that comparison
    gives false positives. Whether a publication's differing n's reflect a
    REAL overlapping-subsample ambiguity (PSYRATS's AHS vs DS) is a judgement
    call made once, by eye, at extraction time -- see the subsample-naming
    addendum -- so it's supplied by the caller instead of guessed here.
    """
    if not ambiguous_publications:
        return False
    present = set(scale_df["publication"].dropna().unique())
    return bool(present & set(ambiguous_publications))


# ---------------------------------------------------------------------------
# Header labels for one block, given the sheet's expanded/single decision.
# ---------------------------------------------------------------------------
def _n_column_label(measure):
    """Column label for a measure's own sample-size column."""
    return "n_total" if measure == TOTAL_COLUMN_LABEL else f"n_{measure}"


def _header_columns(measures, expanded):
    """Return the column labels for one block.

    Single-column layout : ["sample_size", "subsample", <measures...>]
    Expanded layout       : ["subsample", "n_<m1>", <m1>, "n_<m2>", <m2>, ...]
    """
    if not expanded:
        return ["sample_size", "subsample"] + measures
    header = ["subsample"]
    for m in measures:
        header.append(_n_column_label(m))
        header.append(m)
    return header


# ---------------------------------------------------------------------------
# Build one block as a GRID: a list of rows, each row a list of cell values.
# ---------------------------------------------------------------------------
def _block_grid(scale_df, per_sub, per_pub, publication, sample_type, measures, expanded):
    """Return (header_row, data_rows) for one block, as plain lists.

    Single-column layout (expanded=False):
      header_row : ["sample_size", "subsample", <measures...>]
      data row   : [sample_size, subsample, <one value per measure>]
    Expanded layout (expanded=True), used when n varies by measure somewhere
    on this sheet:
      header_row : ["subsample", "n_<m1>", <m1>, "n_<m2>", <m2>, ...]
      data row   : [subsample, <n, value pair per measure>]
    Either way, data_rows is one list per subsample plus a final "total" row;
    empty list if this publication has no data for this sample_type.
    """
    header = _header_columns(measures, expanded)

    block_source = scale_df[
        (scale_df["publication"] == publication)
        & (scale_df["sample_type"] == sample_type)
    ]
    subsamples = sorted(block_source["subsample"].dropna().unique())
    if not subsamples:
        return header, []          # nothing for this sample_type -> stays blank

    data_rows = []
    for subsample in subsamples:
        vals = _measure_values(
            per_sub, measures,
            {"publication": publication, "sample_type": sample_type,
             "subsample": subsample})
        if expanded:
            row = [subsample]
            for measure, value in zip(measures, vals):
                n = _measure_sample_size(scale_df, publication, sample_type,
                                          subsample, measure)
                row.append(n)
                row.append(value)
        else:
            n = _subsample_size(scale_df, publication, sample_type, subsample)
            row = [n, subsample] + vals
        data_rows.append(row)

    # publication total (weighted over this block's subsamples).
    total_vals = _measure_values(
        per_pub, measures,
        {"publication": publication, "sample_type": sample_type})
    if expanded:
        # Sum WITHIN each measure -- summing across measures would
        # double-count participants who contribute to more than one
        # subscale (see the module docstring).
        per_measure_n = _summed_sample_size_per_measure(
            scale_df, sample_type, measures, publication)
        total_row = ["total"]
        for measure, value in zip(measures, total_vals):
            total_row.append(per_measure_n.get(measure))
            total_row.append(value)
    else:
        # The total row reports the SUMMED n (each subsample counted once).
        total_n = _summed_sample_size(scale_df, sample_type, publication)
        total_row = [total_n, "total"] + total_vals
    data_rows.append(total_row)

    return header, data_rows


# ---------------------------------------------------------------------------
# Pad the shorter block so the two "total" rows line up on the same row.
# ---------------------------------------------------------------------------
def _pad_blocks_to_align_totals(hc_rows, pt_rows, row_width):
    """Pad the shorter block with blank rows so both totals end on the same line.

    Each block's rows are [subsample rows..., total row] (or empty if the block
    has no data). We separate the subsample rows from the final total row, pad
    the shorter subsample list with all-blank rows so both have the same number
    of subsample rows, then re-append each block's total row. An empty block is
    padded to the other block's full height (subsamples + total) and stays blank.

    `row_width` is the number of columns in one block's header (differs
    between the single-column and expanded layouts -- see `_header_columns`).
    """
    blank = [None] * row_width

    def split(rows):
        # returns (subsample_rows, total_row_or_None)
        if not rows:
            return [], None
        return rows[:-1], rows[-1]

    hc_subs, hc_total = split(hc_rows)
    pt_subs, pt_total = split(pt_rows)

    # pad the shorter subsample list up to the longer one
    n = max(len(hc_subs), len(pt_subs))
    hc_subs = hc_subs + [list(blank) for _ in range(n - len(hc_subs))]
    pt_subs = pt_subs + [list(blank) for _ in range(n - len(pt_subs))]

    # reassemble; a block that had no data gets an all-blank total line too,
    # so the two blocks stay the same height and the sheet lines up
    hc_out = hc_subs + [hc_total if hc_total is not None else list(blank)]
    pt_out = pt_subs + [pt_total if pt_total is not None else list(blank)]

    # if BOTH blocks were empty, return empty so nothing is written
    if hc_total is None and pt_total is None and n == 0:
        return [], []
    return hc_out, pt_out


# ---------------------------------------------------------------------------
# Write a grid (header + data rows) at an explicit top-left cell.
# ---------------------------------------------------------------------------
def _write_grid(ws, header, data_rows, top_row, left_col):
    """Write one block's header and data at explicit (row, col) coordinates.

    Rows/cols are 1-based (openpyxl convention). If data_rows is empty, only
    the header is written (the block is otherwise blank). Returns the number of
    rows written (header + data), so the caller can advance past the block.
    """
    # header
    for j, label in enumerate(header):
        ws.cell(row=top_row, column=left_col + j, value=label)
    # data
    for i, row_values in enumerate(data_rows, start=1):
        for j, value in enumerate(row_values):
            # None stays an empty cell; that is how blanks are represented
            if value is not None:
                ws.cell(row=top_row + i, column=left_col + j, value=value)
    return 1 + len(data_rows)


# ---------------------------------------------------------------------------
# Write one scale's worksheet, stacked publication by publication.
# ---------------------------------------------------------------------------
def _write_scale_sheet(ws, scale, scale_df, measures, per_sub, per_pub, grand,
                        ambiguous_publications):
    """Write the whole worksheet for one scale, using explicit coordinates."""
    # Decide ONCE for the whole sheet whether this scale needs the expanded,
    # per-measure sample-size columns -- see _needs_expanded_columns / module
    # docstring. This is a caller-supplied list, not inferred from the data.
    expanded = _needs_expanded_columns(scale_df, ambiguous_publications)
    block_width = len(_header_columns(measures, expanded))
    gap = 1                                 # blank column between the two blocks
    # 1-based left column of each block
    left_col = {
        "healthy_controls": 1,
        "patients": 1 + block_width + gap,
    }

    row = 1  # current 1-based row cursor

    # title
    ws.cell(row=row, column=1,
            value=f"{scale} — weighted mean values "
                  f"(Option B; redundant aggregates excluded)")
    row += 2

    for publication in sorted(scale_df["publication"].dropna().unique()):
        # publication name
        ws.cell(row=row, column=1, value=publication)
        row += 1

        # sample_type header line: label each block region
        ws.cell(row=row, column=1, value="sample_type")
        ws.cell(row=row, column=left_col["healthy_controls"] + 1,
                value="healthy_controls")
        ws.cell(row=row, column=left_col["patients"], value="patients")
        row += 1

        # build both grids
        hc_header, hc_rows = _block_grid(
            scale_df, per_sub, per_pub, publication, "healthy_controls", measures, expanded)
        pt_header, pt_rows = _block_grid(
            scale_df, per_sub, per_pub, publication, "patients", measures, expanded)

        # Align the two "total" rows on the same line. In each block the LAST
        # data row is the total; the rows before it are the subsamples. We pad
        # the block that has fewer subsamples with blank rows so both totals end
        # up on the same line. (An empty block has no rows and stays blank.)
        hc_rows, pt_rows = _pad_blocks_to_align_totals(hc_rows, pt_rows, block_width)

        # write both at the SAME top row, each in its own column region
        hc_h = _write_grid(ws, hc_header, hc_rows, row, left_col["healthy_controls"])
        pt_h = _write_grid(ws, pt_header, pt_rows, row, left_col["patients"])

        # advance past the taller block, then a blank spacer row
        row += max(hc_h, pt_h) + 2

    # grand total section
    ws.cell(row=row, column=1, value="GRAND TOTAL (all publications)")
    row += 1
    ws.cell(row=row, column=1, value="sample_type")
    ws.cell(row=row, column=left_col["healthy_controls"] + 1,
            value="healthy_controls")
    ws.cell(row=row, column=left_col["patients"], value="patients")
    row += 1

    for sample_type in BLOCK_ORDER:
        header = _header_columns(measures, expanded)
        vals = _measure_values(grand, measures, {"sample_type": sample_type})
        # only write a grand-total row if there is at least one value
        if all(v is None for v in vals):
            data_rows = []
        elif expanded:
            # sum WITHIN each measure over ALL publications -- same reasoning
            # as the per-publication total row above.
            per_measure_n = _summed_sample_size_per_measure(
                scale_df, sample_type, measures, publication=None)
            grand_row = ["grand total"]
            for measure, value in zip(measures, vals):
                grand_row.append(per_measure_n.get(measure))
                grand_row.append(value)
            data_rows = [grand_row]
        else:
            # grand total reports summed n over ALL publications (each
            # subsample counted once)
            grand_n = _summed_sample_size(scale_df, sample_type, publication=None)
            data_rows = [[grand_n, "grand total"] + vals]
        _write_grid(ws, header, data_rows, row, left_col[sample_type])


# ---------------------------------------------------------------------------
# Top-level: build all scales and write the workbook.
# ---------------------------------------------------------------------------
def write_presentation_excel(df, output_path, subscale_orders=None,
                              ambiguous_publications=None):
    """Write the presentation workbook: one sheet per scale.

    subscale_orders : optional dict {scale: [subscale, ...]} to fix column order.
    ambiguous_publications : optional list/set of publication names known to
        have overlapping/non-additive subsample sample sizes across subscales
        (e.g. PSYRATS's Favrod et al. 2012, Woodward et al. 2014 -- see the
        subsample-naming addendum). Any scale sheet containing one of these
        publications switches its ENTIRE sheet to a per-measure sample-size
        column layout; see the module docstring. This is not inferred from
        the data -- pass it in explicitly once you've identified them.
    """
    subscale_orders = subscale_orders or {}
    ambiguous_publications = ambiguous_publications or []

    # Pre-compute the three aggregation levels once, reusing aggregate.py.
    per_sub = aggregate.aggregate(
        df, "mean",
        group_cols=["scale", "publication", "sample_type", "subsample",
                    "subscale", "record_type"])
    per_pub = aggregate.aggregate(
        df, "mean",
        group_cols=["scale", "publication", "sample_type",
                    "subscale", "record_type"])
    grand = aggregate.aggregate(
        df, "mean",
        group_cols=["scale", "sample_type", "subscale", "record_type"])

    wb = Workbook()
    wb.remove(wb.active)   # drop the default empty sheet

    for scale in sorted(df["scale"].dropna().unique()):
        scale_df = df[df["scale"] == scale]
        measures = _measure_columns(scale_df, subscale_orders.get(scale))
        ws = wb.create_sheet(title=scale[:31])   # 31-char sheet-name limit
        _write_scale_sheet(
            ws, scale, scale_df, measures,
            per_sub[per_sub["scale"] == scale],
            per_pub[per_pub["scale"] == scale],
            grand[grand["scale"] == scale],
            ambiguous_publications)

    wb.save(output_path)
    print(f"wrote presentation Excel -> {output_path}")
